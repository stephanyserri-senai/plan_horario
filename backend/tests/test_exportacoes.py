"""Exportações prontas: horário das turmas, agenda dos instrutores, uso das salas e cadastros no formato de importação."""
import io

from openpyxl import load_workbook
from sqlalchemy import select

from app import demo
from app.services import exportacao_planilhas, importacao
from tests.test_otimizacao import executar
from app.services.modelos_planilha import gerar_modelo


def test_exportacoes_planilhas(db, mundo):
    assert executar(db, mundo).status == "concluida"
    cen = mundo.cen
    from app.services import indicadores
    d0, d1 = indicadores.periodo_padrao(db, cen)
    for arq, primeira in ((exportacao_planilhas.horario_turmas(db, cen.id), "Resumo"),
                          (exportacao_planilhas.agenda_instrutores(db, cen.id, d0, d1), "Ocupação"),
                          (exportacao_planilhas.uso_ambientes(db, cen.id, d0, d1), "Ocupação")):
        wb = load_workbook(io.BytesIO(arq))
        assert wb.sheetnames[0] == primeira and len(wb.sheetnames) > 1
    wb = load_workbook(io.BytesIO(exportacao_planilhas.horario_turmas(db, cen.id)))
    assert mundo.turma1.codigo in wb.sheetnames if hasattr(mundo, "turma1") else True
    assert any(ws.max_row > 5 for ws in wb.worksheets[1:])


def test_cadastro_exportado_reimporta_sem_alteracao(db, mundo):
    demo.carregar(db)
    for tipo in ("estrutura_curricular", "ambientes", "levantamento_instrutores", "turmas"):
        linhas = [l for l in exportacao_planilhas.linhas_cadastro(db, tipo) if str(l[0]).startswith("DEMO-")]  # só a demonstração (completa)
        assert linhas, tipo
        imp = importacao.gerar_previa(db, mundo.ctx, f"{tipo}.xlsx", gerar_modelo(db, tipo, linhas))
        assert imp.tipo == tipo
        acoes = {i["acao"] for i in imp.previa["itens"]}
        assert acoes == {"sem_alteracao"}, [i for i in imp.previa["itens"] if i["acao"] == "erro"][:3]
        db.rollback()
