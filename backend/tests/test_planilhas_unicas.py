"""Modelos de planilha + importação em planilha única (cursos/matrizes/UCs, salas, levantamento de instrutores, turmas)."""
import io

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.estrutura import garantir_estrutura_fixa
from app.models import (Ambiente, CompetenciaInstrutor, Curso, DisponibilidadeInstrutor, HabilitacaoInstrutor, Instrutor, MatrizCurricular,
                        TipoAmbiente, Turma, TurmaUC, UCPrerequisito, UnidadeCurricular)
from app.services import importacao
from app.services.modelos_planilha import gerar_modelo


def preencher(db, tipo, linhas):
    """Abre o modelo gerado e escreve as linhas na área amarela (logo após os exemplos), como o usuário faria."""
    wb = load_workbook(io.BytesIO(gerar_modelo(db, tipo)))
    ws = wb["Preencher"]
    inicio = next(i for i in range(4, 20) if not str(ws.cell(row=i, column=1).value or "").startswith("EXEMPLO"))
    for k, linha in enumerate(linhas):
        for j, v in enumerate(linha, start=1):
            ws.cell(row=inicio + k, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def importar(db, ctx, tipo, conteudo):
    imp = importacao.gerar_previa(db, ctx, f"{tipo}.xlsx", conteudo)
    assert imp.tipo == tipo, f"tipo detectado {imp.tipo}"  # reconhecido pelos cabeçalhos do modelo
    erros = [i for i in imp.previa["itens"] if i["acao"] == "erro"]
    assert not erros, erros
    rel = importacao.confirmar(db, ctx, imp.id)
    db.commit()
    return imp, rel


def test_modelos_e_importacao_em_planilha_unica(db, mundo):
    m = mundo
    garantir_estrutura_fixa(db)
    db.commit()
    # 1) estrutura curricular: 1 linha por UC; cria curso, matriz, UCs, tipos de sala e pré-requisito
    est = preencher(db, "estrutura_curricular", [
        ["MEC-TEC", "Técnico em Mecânica", "MECÂNICA", "presencial", "técnico", 1200, "2026.1", "DTM", "Desenho Técnico Mecânico", 40, 1, "Sala de aula", 1, "", ""],
        ["MEC-TEC", "Técnico em Mecânica", "MECÂNICA", "presencial", "técnico", 1200, "2026.1", "USI", "Usinagem", 60, 2, "Oficina mecânica", 2, "DTM", "NR-12"],
    ])
    imp, rel = importar(db, m.ctx, "estrutura_curricular", est)
    assert len(imp.previa["itens"]) == 2  # as linhas EXEMPLO foram ignoradas
    curso = db.scalar(select(Curso).where(Curso.codigo == "MEC-TEC"))
    assert curso and curso.area_id and rel["cursos"] == 1 and rel["matrizes"] == 1 and rel["criados"] == 2
    usi = db.scalar(select(UnidadeCurricular).where(UnidadeCurricular.codigo == "USI"))
    dtm = db.scalar(select(UnidadeCurricular).where(UnidadeCurricular.codigo == "DTM"))
    assert usi.encontros_semanais == 2 and usi.certificacao_exigida == "NR-12" and usi.duracao_bloco_periodos == 4
    assert [t.nome for t in usi.tipos_ambiente] == ["Oficina mecânica"]
    assert db.get(UCPrerequisito, (usi.id, dtm.id))
    # reimportar o mesmo arquivo não duplica
    imp2 = importacao.gerar_previa(db, m.ctx, "estrutura_curricular.xlsx", est)
    assert {i["acao"] for i in imp2.previa["itens"]} == {"sem_alteracao"}
    db.rollback()

    # 2) salas e laboratórios: tipo novo criado; TODAS as áreas e várias áreas
    salas = preencher(db, "ambientes", [
        ["BM-S201", "Sala 201", "BEIRA-MAR", "Sala de aula", "Não", "TODAS", 40, 36, "Bloco A", "Sim", 0],
        ["BM-OF1", "Oficina de Usinagem", "BEIRA-MAR", "Oficina mecânica", "Sim", "MECÂNICA; AUTOMOTIVA", 20, 16, "Bloco C", "Não", 15],
    ])
    importar(db, m.ctx, "ambientes", salas)
    of1 = db.scalar(select(Ambiente).where(Ambiente.codigo == "BM-OF1"))
    s201 = db.scalar(select(Ambiente).where(Ambiente.codigo == "BM-S201"))
    assert s201.todas_areas and of1.area_id and db.get(TipoAmbiente, of1.tipo_ambiente_id).especializado
    assert db.scalar(select(func.count()).select_from(TipoAmbiente).where(TipoAmbiente.nome == "Oficina mecânica")) == 1

    # 3) levantamento de instrutores: áreas, turnos × dias, competências e certificação
    lev = preencher(db, "levantamento_instrutores", [
        ["M-100", "Paulo Torres", "paulo@x.br", "BEIRA-MAR", "MECÂNICA; AUTOMOTIVA", 40, 28, 8, "Manhã; Tarde", "seg a qui", "DTM; USI", "", "NR-12",
         "Técnico em mecânica; 10 anos em usinagem CNC"],
    ])
    importar(db, m.ctx, "levantamento_instrutores", lev)
    paulo = db.scalar(select(Instrutor).where(Instrutor.matricula == "M-100"))
    disp = db.scalars(select(DisponibilidadeInstrutor).where(DisponibilidadeInstrutor.instrutor_id == paulo.id)).all()
    assert len(disp) == 8 and {d.dia_semana for d in disp} == {0, 1, 2, 3}  # 2 turnos × 4 dias
    assert {d.hora_inicio.strftime("%H:%M") for d in disp} == {"07:30", "13:30"}
    assert db.scalar(select(func.count()).select_from(CompetenciaInstrutor).where(CompetenciaInstrutor.instrutor_id == paulo.id)) == 2
    assert db.scalars(select(HabilitacaoInstrutor).where(HabilitacaoInstrutor.instrutor_id == paulo.id)).first().nome == "NR-12"
    assert paulo.areas_competencia.startswith("Técnico em mecânica")

    # 4) turmas: UCs da matriz ligadas automaticamente
    tur = preencher(db, "turmas", [["MEC-1", "Mecânica 2026/2 Manhã", "MEC-TEC", "2026.1", "BEIRA-MAR", "Manhã", 28, "03/08/2026", "18/12/2026", "seg a sex", ""]])
    imp, rel = importar(db, m.ctx, "turmas", tur)
    t = db.scalar(select(Turma).where(Turma.codigo == "MEC-1"))
    assert t.area_id == curso.area_id
    assert db.scalar(select(func.count()).select_from(TurmaUC).where(TurmaUC.turma_id == t.id)) == 2


def test_modelo_tem_instrucoes_e_listas(db, mundo):
    garantir_estrutura_fixa(db)
    db.commit()
    wb = load_workbook(io.BytesIO(gerar_modelo(db, "levantamento_instrutores")))
    assert wb.sheetnames == ["Preencher", "Instruções", "Listas"] or set(wb.sheetnames) == {"Preencher", "Instruções", "Listas"}
    listas = wb["Listas"]
    valores = [c.value for c in listas["A"] if c.value]
    assert "BEIRA-MAR" in valores and "PORTO" in valores
    assert wb["Preencher"].data_validations.dataValidation  # listas de seleção
