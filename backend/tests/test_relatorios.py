"""Relatórios: todos baixam em Excel e PDF, com cabeçalhos legíveis; grade da semana com uma aba por turma; pacote ZIP."""
import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.conftest import SEG
from tests.test_otimizacao import executar


def test_relatorios_excel_pdf_e_pacote(db, mundo):
    m = mundo
    assert executar(db, m).status == "concluida"
    c = TestClient(app)
    h = {"Authorization": "Bearer " + c.post("/api/auth/login", json={"email": "admin@teste", "senha": "senha-forte-1"}).json()["token"]}
    cid = m.cen.id
    urls = [f"/api/exportar/grade?cenario_id={cid}", f"/api/exportar/grade?cenario_id={cid}&filtro=turma&filtro_id={m.t1.id}",
            f"/api/exportar/grade?cenario_id={cid}&semana={SEG}", f"/api/exportar/horario_turmas?cenario_id={cid}",
            f"/api/exportar/agenda_instrutores?cenario_id={cid}", f"/api/exportar/uso_ambientes?cenario_id={cid}",
            *[f"/api/exportar/indicadores/{t}?cenario_id={cid}" for t in ("ambientes", "instrutores", "turmas", "gargalos")]]
    for u in urls:
        for f in ("xlsx", "pdf"):
            r = c.get(f"{u}&formato={f}", headers=h)
            assert r.status_code == 200, (u, f, r.text[:300])
            assert r.content[:4] == (b"PK\x03\x04" if f == "xlsx" else b"%PDF"), (u, f)
    # cabeçalhos em português, com identificação do semestre
    wb = load_workbook(io.BytesIO(c.get(f"/api/exportar/indicadores/turmas?cenario_id={cid}&formato=xlsx", headers=h).content))
    ws = wb.active
    textos = [str(v) for row in ws.iter_rows(max_row=10, values_only=True) for v in row if v]
    assert "Carga horária por turma e UC" in textos and any(t.startswith("Semestre:") for t in textos)
    assert {"Turma", "Unidade curricular", "Falta (h)", "Risco de atraso"} <= set(textos)
    assert not any("_" in t and t.islower() for t in textos), textos  # nada de nome técnico
    # grade da semana: visão geral + uma aba por turma + lista
    wb = load_workbook(io.BytesIO(c.get(f"/api/exportar/grade?cenario_id={cid}&semana={SEG}&formato=xlsx", headers=h).content))
    assert wb.sheetnames[0] == "Semana (todas as turmas)" and wb.sheetnames[-1] == "Aulas da semana"
    assert {m.t1.codigo, m.t2.codigo} <= set(wb.sheetnames)
    # todos de uma vez
    r = c.get(f"/api/exportar/pacote?cenario_id={cid}&formato=ambos", headers=h)
    assert r.status_code == 200
    nomes = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert len(nomes) == 16 and {n.rsplit(".", 1)[1] for n in nomes} == {"xlsx", "pdf"}
