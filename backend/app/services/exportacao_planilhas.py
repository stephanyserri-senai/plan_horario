"""Exportações prontas para uso (XLSX formatado).

- Horário das turmas: uma aba por turma, semana a semana (linhas) × dias (colunas),
  com a UC, o instrutor e a sala de cada dia, e o resumo de carga horária das UCs.
- Agenda dos instrutores: aba "Ocupação" (resumo) e uma aba por instrutor.
- Uso das salas e laboratórios: aba "Ocupação" (resumo) e uma aba por ambiente.
- Cadastros no mesmo formato dos modelos de importação (edite e reenvie).
"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (Ambiente, AmbienteArea, Area, CompetenciaInstrutor, Curso, DisponibilidadeInstrutor, HabilitacaoInstrutor, Instrutor,
                        InstrutorArea, MatrizCurricular, TipoAmbiente, Turma, TurmaUC, Turno, UCPrerequisito, Unidade, UnidadeCurricular)
from app.services import exportacao, indicadores
from app.services.planejamento import situacao_carga_horaria

FONTE = "Arial"
AZUL = PatternFill("solid", fgColor="1F4E79")
CLARO = PatternFill("solid", fgColor="F2F6FA")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
PALETA = ["DCE9F7", "E2F0D9", "FFF2CC", "FCE4D6", "E4DFEC", "DDEBF7", "D9F2EF", "F8E1E7", "EDEDED", "FFE699"]
DIAS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado"]


def _nome_aba(nome: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", nome)[:31] or "aba"
    n, k = base, 2
    while n.lower() in usados:
        sufixo = f" ({k})"
        n, k = base[: 31 - len(sufixo)] + sufixo, k + 1
    usados.add(n.lower())
    return n


def _cab(ws, linha: int, titulos: list[str], larguras: list[int] | None = None):
    for j, t in enumerate(titulos, 1):
        c = ws.cell(linha, j, t)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF")
        c.fill = AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
        if larguras:
            ws.column_dimensions[get_column_letter(j)].width = larguras[j - 1]


def _titulo(ws, titulo: str, subtitulo: str | None = None):
    ws["A1"] = titulo
    ws["A1"].font = Font(name=FONTE, bold=True, size=14, color="1F4E79")
    ws.row_dimensions[1].height = 22
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(name=FONTE, size=9, italic=True, color="595959")


def _tabela(ws, linha: int, cols: list[tuple[str, str, int]], dados: list[dict]):
    """cols: (chave, título, largura). Retorna a próxima linha livre."""
    _cab(ws, linha, [t for _, t, _ in cols], [w for _, _, w in cols])
    for i, d in enumerate(dados, linha + 1):
        for j, (k, _, _) in enumerate(cols, 1):
            c = ws.cell(i, j, d.get(k))
            c.font = Font(name=FONTE, size=10)
            c.border = BORDA
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if i % 2 == 0:
                c.fill = CLARO
    return linha + len(dados) + 1


def _salvar(wb: Workbook) -> bytes:
    from openpyxl.worksheet.properties import PageSetupProperties
    if len(wb.sheetnames) > 1 and not wb.worksheets[0].max_row:
        wb.remove(wb.worksheets[0])
    for ws in wb.worksheets:  # leitura e impressão: sem linhas de grade, paisagem, cabe na largura, página numerada
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.oddFooter.center.text = "Página &P de &N"
        ws.oddFooter.right.text = exportacao.INSTITUICAO
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _semanas(linhas: list[dict]) -> list[date]:
    if not linhas:
        return []
    ini = min(l["_data"] for l in linhas)
    fim = max(l["_data"] for l in linhas)
    seg = ini - timedelta(days=ini.weekday())
    res = []
    while seg <= fim:
        res.append(seg)
        seg += timedelta(days=7)
    return res


def _quadro(ws, linha: int, linhas: list[dict], texto, chave_cor) -> int:
    """Quadro semana × dia. `texto(lista_de_aulas_do_dia)` -> str; `chave_cor(aula)` agrupa cores."""
    usa_sab = any(l["_data"].weekday() == 5 for l in linhas)
    dias = DIAS if usa_sab else DIAS[:5]
    _cab(ws, linha, ["Semana"] + dias, [16] + [30] * len(dias))
    por_dia = defaultdict(list)
    for l in linhas:
        por_dia[l["_data"]].append(l)
    cores: dict[str, PatternFill] = {}
    r = linha + 1
    for seg in _semanas(linhas):
        c = ws.cell(r, 1, f"{seg:%d/%m} a {seg + timedelta(days=len(dias) - 1):%d/%m}")
        c.font = Font(name=FONTE, bold=True, size=10)
        c.alignment = Alignment(vertical="center")
        c.border = BORDA
        maior = 1
        for j in range(len(dias)):
            d = seg + timedelta(days=j)
            aulas = sorted(por_dia.get(d, []), key=lambda a: a["_ini"])
            cel = ws.cell(r, j + 2)
            cel.border = BORDA
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(name=FONTE, size=9)
            if aulas:
                txt = texto(aulas)
                cel.value = txt
                k = chave_cor(aulas[0])
                cores.setdefault(k, PatternFill("solid", fgColor=PALETA[len(cores) % len(PALETA)]))
                cel.fill = cores[k]
                maior = max(maior, txt.count("\n") + 1)
        ws.row_dimensions[r].height = max(15, 12.5 * maior)
        r += 1
    ws.freeze_panes = ws.cell(linha + 1, 2)
    return r + 1


def _blocos_dia(aulas: list[dict]) -> list[tuple[dict, str]]:
    """Blocos do dia (mesma turma/UC): (primeira aula, horário total)."""
    blocos: dict[tuple, list[dict]] = {}
    for a in aulas:
        blocos.setdefault((a["turno"], a["turma"], a["uc"]), []).append(a)
    return [(bs[0], f"{min(b['inicio'] for b in bs)}–{max(b['fim'] for b in bs)}") for bs in blocos.values()]


def _agrupar_dia(aulas: list[dict], rotulo) -> list[str]:
    """Junta os períodos de um mesmo bloco (mesma turma/UC no dia) numa linha com o horário total."""
    blocos: dict[tuple, list[dict]] = {}
    for a in aulas:
        blocos.setdefault((a["turno"], a["turma"], a["uc"]), []).append(a)
    res = []
    for _k, bs in blocos.items():
        ini, fim = min(b["inicio"] for b in bs), max(b["fim"] for b in bs)
        res.append(rotulo(bs[0], f"{ini}–{fim}"))
    return res


# ============================================================================ horário das turmas
def horario_turmas(db: Session, cenario_id: int, turma_id: int | None = None, unidades=None) -> bytes:
    linhas = exportacao.linhas_grade(db, cenario_id, "turma" if turma_id else None, turma_id, unidades_permitidas=unidades)
    q = select(Turma).where(Turma.ativo.is_(True), Turma.simulacao_id.is_(None)).order_by(Turma.codigo)
    if turma_id:
        q = q.where(Turma.id == turma_id)
    if unidades is not None:
        q = q.where(Turma.unidade_id.in_(unidades))
    turmas = db.scalars(q).all()
    cursos = {c.id: c for c in db.scalars(select(Curso))}
    turnos = {t.id: t.nome for t in db.scalars(select(Turno))}
    nome_cen = exportacao.nome_cenario(db, cenario_id)
    por_turma = defaultdict(list)
    for l in linhas:
        por_turma[l["turma"]].append(l)
    ead_turma = defaultdict(list)  # turmas 80-20: o dia EAD aparece no quadro da turma
    if turmas:
        for l in exportacao.linhas_ead(db, min(t.data_inicio for t in turmas), max(t.data_fim for t in turmas), "turma" if turma_id else None,
                                       turma_id, unidades):
            ead_turma[l["turma"]].append(l)
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    _titulo(resumo, "Horário das turmas", f"{nome_cen} · cada turma tem sua aba, com a semana a semana e a carga horária das UCs")
    usados = {"resumo"}
    linhas_resumo = []
    for t in turmas:
        tucs = db.scalars(select(TurmaUC).where(TurmaUC.turma_id == t.id, TurmaUC.ativo.is_(True)).order_by(TurmaUC.id)).all()
        infos = [situacao_carga_horaria(db, cenario_id, x) for x in tucs]
        tot = sum(i["total_h"] for i in infos)
        com = sum(i["realizada_h"] + i["programada_h"] for i in infos)
        aulas_t = por_turma.get(t.codigo, [])
        linhas_resumo.append({"turma": t.codigo, "nome": t.nome, "curso": cursos[t.curso_id].nome if t.curso_id in cursos else "",
                              "turno": turnos.get(t.turno_id, ""), "periodo": f"{t.data_inicio:%d/%m/%Y} a {t.data_fim:%d/%m/%Y}",
                              "ucs": len(infos), "total": round(tot, 1), "com": round(com, 1), "falta": round(max(tot - com, 0), 1),
                              "pct": f"{round(100 * com / tot) if tot else 0}%",
                              "termino": max(a["_data"] for a in aulas_t).strftime("%d/%m/%Y") if aulas_t else "—"})
        ws = wb.create_sheet(_nome_aba(t.codigo, usados))
        _titulo(ws, f"{t.codigo} — {t.nome}", f"{cursos[t.curso_id].nome if t.curso_id in cursos else ''} · {turnos.get(t.turno_id, '')} · "
                f"{t.data_inicio:%d/%m/%Y} a {t.data_fim:%d/%m/%Y} · {nome_cen}")
        if aulas_t:
            com_ead = aulas_t + [e for e in ead_turma.get(t.codigo, []) if min(a["_data"] for a in aulas_t) <= e["_data"] <= max(a["_data"] for a in aulas_t)]
            r = _quadro(ws, 4, com_ead, lambda aulas: "\n".join(
                "EAD – atividades a distância" if a.get("_ead") else f"{a['uc']}\n{a['instrutor']} · {a['ambiente']} · {h}"
                for a, h in ((bs, h) for bs, h in _blocos_dia(aulas))), lambda a: "EAD" if a.get("_ead") else a["uc"])
        else:
            ws["A4"] = "Ainda não há aulas na grade para esta turma."
            r = 6
        ws.cell(r, 1, "Unidades curriculares").font = Font(name=FONTE, bold=True, size=11)
        nomes_inst = defaultdict(set)
        for a in aulas_t:
            nomes_inst[a["uc"].split(" – ")[0]].add(a["instrutor"])
        _tabela(ws, r + 1, [("uc", "UC", 16), ("nome", "Nome", 30), ("total", "Carga (h)", 30), ("com", "Com horário (h)", 30),
                            ("falta", "Falta (h)", 30), ("instrutores", "Instrutor(es)", 30)],
                [{"uc": i["uc_codigo"], "nome": i["uc_nome"], "total": i["total_h"], "com": round(i["realizada_h"] + i["programada_h"], 1),
                  "falta": i["restante_h"], "instrutores": ", ".join(sorted(nomes_inst.get(i["uc_codigo"], [])))} for i in infos])
    _tabela(resumo, 4, [("turma", "Turma", 16), ("nome", "Nome", 32), ("curso", "Curso", 30), ("turno", "Turno", 10), ("periodo", "Período", 24),
                        ("ucs", "UCs", 7), ("total", "Carga total (h)", 12), ("com", "Com horário (h)", 12), ("falta", "Falta (h)", 10),
                        ("pct", "% com horário", 10), ("termino", "Última aula", 12)], linhas_resumo)
    resumo.freeze_panes = "A5"
    return _salvar(wb)


# ============================================================================ agenda dos instrutores
def agenda_instrutores(db: Session, cenario_id: int, d0: date, d1: date, instrutor_id: int | None = None, unidades=None,
                       mostrar_pessoais: bool = False) -> bytes:
    linhas = exportacao.linhas_grade(db, cenario_id, "instrutor" if instrutor_id else None, instrutor_id, unidades_permitidas=unidades)
    oc = indicadores.ocupacao_instrutores(db, cenario_id, d0, d1, unidades_permitidas=unidades, mostrar_pessoais=mostrar_pessoais)["instrutores"]
    if instrutor_id:
        oc = [o for o in oc if o["instrutor_id"] == instrutor_id]
    nome_cen = exportacao.nome_cenario(db, cenario_id)
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Ocupação"
    _titulo(resumo, "Ocupação dos instrutores", f"{nome_cen} · {d0:%d/%m/%Y} a {d1:%d/%m/%Y} · ocupação = horas de aula por semana (média) ÷ capacidade para ensino")
    fim = _tabela(resumo, 4, [("nome", "Instrutor", 30), ("matricula", "Matrícula", 14), ("carga_contratual_h", "Contrato (h/sem)", 12),
                              ("capacidade_efetiva_h", "Capacidade p/ aulas (h/sem)", 14), ("alocada_media_semanal_h", "Aulas (h/sem, média)", 14),
                              ("maior_semana_h", "Semana mais cheia (h)", 14), ("restante_semanal_h", "Folga (h/sem)", 12),
                              ("ocupacao_pct", "Ocupação (%)", 12), ("alocada_total_h", "Total no período (h)", 14)],
                  sorted(oc, key=lambda o: -o["ocupacao_pct"]))
    for i in range(5, fim):
        c = resumo.cell(i, 8)
        v = c.value or 0
        c.fill = PatternFill("solid", fgColor="FDE7E5" if v > 100 else "FFF3D6" if v >= 85 else "E3F4EA" if v >= 40 else "F2F2F2")
    resumo.freeze_panes = "A5"
    usados = {"ocupação"}
    por_inst = defaultdict(list)
    for l in linhas:
        por_inst[l["instrutor"]].append(l)
    for o in sorted(oc, key=lambda o: o["nome"]):
        ws = wb.create_sheet(_nome_aba(o["nome"], usados))
        _titulo(ws, f"Agenda — {o['nome']}", f"Ocupação {o['ocupacao_pct']}% · {o['alocada_media_semanal_h']} h/semana em média · {nome_cen}")
        aulas = por_inst.get(o["nome"], [])
        if aulas:
            _quadro(ws, 4, aulas, lambda aulas: "\n".join(_agrupar_dia(aulas, lambda a, h: f"{a['turno']} {h}\n{a['turma']} · {a['uc'].split(' – ')[0]} · {a['ambiente']}")),
                    lambda a: a["turma"])
        else:
            ws["A4"] = "Sem aulas atribuídas no semestre."
    return _salvar(wb)


# ============================================================================ salas e laboratórios
def uso_ambientes(db: Session, cenario_id: int, d0: date, d1: date, ambiente_id: int | None = None, unidades=None) -> bytes:
    linhas = exportacao.linhas_grade(db, cenario_id, "ambiente" if ambiente_id else None, ambiente_id, unidades_permitidas=unidades)
    oc = indicadores.ocupacao_ambientes(db, cenario_id, d0, d1, unidades_permitidas=unidades)["ambientes"]
    ambs = {a.id: a for a in db.scalars(select(Ambiente))}
    if ambiente_id:
        oc = [o for o in oc if o.get("ambiente_id") == ambiente_id]
    nome_cen = exportacao.nome_cenario(db, cenario_id)
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Ocupação"
    _titulo(resumo, "Ocupação das salas e laboratórios", f"{nome_cen} · {d0:%d/%m/%Y} a {d1:%d/%m/%Y}")
    CLASSE = {"subutilizado": "Pouco usada", "adequado": "Adequada", "demanda_excessiva": "Sobrecarregada", "sem_disponibilidade": "Sem horário disponível"}
    fim = _tabela(resumo, 4, [("codigo", "Código", 14), ("nome", "Sala / laboratório", 30), ("tipo", "Tipo", 24), ("capacidade_fisica", "Lugares", 9),
                              ("aulas", "Aulas", 8), ("horas_disponiveis", "Horas disponíveis", 12), ("horas_ocupadas", "Horas ocupadas", 12),
                              ("horas_ociosas", "Horas livres", 12), ("taxa_ocupacao_pct", "Ocupação (%)", 11), ("situacao_txt", "Situação", 20)],
                  [{**o, "situacao_txt": CLASSE.get(o["classificacao"], o["classificacao"])} for o in sorted(oc, key=lambda o: -o["taxa_ocupacao_pct"])])
    for i in range(5, fim):
        v = resumo.cell(i, 9).value or 0
        resumo.cell(i, 9).fill = PatternFill("solid", fgColor="FDE7E5" if v > 90 else "E3F4EA" if v >= 40 else "FFF3D6")
    resumo.freeze_panes = "A5"
    usados = {"ocupação"}
    por_amb = defaultdict(list)
    for l in linhas:
        por_amb[l["ambiente"]].append(l)
    lista = [ambs[o["ambiente_id"]] for o in oc if o["ambiente_id"] in ambs]
    for a in sorted(lista, key=lambda a: a.codigo):
        ws = wb.create_sheet(_nome_aba(a.codigo, usados))
        _titulo(ws, f"{a.codigo} — {a.nome}", f"Capacidade {a.capacidade_fisica} lugares · {nome_cen}")
        aulas = por_amb.get(a.codigo, [])
        if aulas:
            _quadro(ws, 4, aulas, lambda aulas: "\n".join(_agrupar_dia(aulas, lambda x, h: f"{x['turno']} {h}\n{x['turma']} · {x['uc'].split(' – ')[0]} · {x['instrutor']}")),
                    lambda x: x["turno"])
        else:
            ws["A4"] = "Sem aulas neste ambiente no semestre."
    return _salvar(wb)


# ============================================================================ cadastros no formato de importação
def _areas_de(db, principal_id, extras_model, campo, obj_id, todas: bool, areas: dict) -> str:
    if todas:
        return "TODAS"
    ids = [principal_id] if principal_id else []
    ids += [x for x in db.scalars(select(extras_model.area_id).where(getattr(extras_model, campo) == obj_id)) if x not in ids]
    return "; ".join(areas[i] for i in ids if i in areas)


def linhas_cadastro(db: Session, tipo: str, unidades=None) -> list[list]:
    """Linhas na mesma ordem de colunas de modelos_planilha.MODELOS[tipo]."""
    areas = {a.id: a.nome for a in db.scalars(select(Area))}
    unid = {u.id: u.codigo for u in db.scalars(select(Unidade))}
    turnos_nome = {t.id: t.nome for t in db.scalars(select(Turno))}
    res = []
    if tipo == "estrutura_curricular":
        for c in db.scalars(select(Curso).where(Curso.ativo.is_(True)).order_by(Curso.codigo)):
            if unidades is not None and c.unidade_id and c.unidade_id not in unidades:
                continue
            for m in db.scalars(select(MatrizCurricular).where(MatrizCurricular.curso_id == c.id, MatrizCurricular.ativo.is_(True)).order_by(MatrizCurricular.versao)):
                ucs = db.scalars(select(UnidadeCurricular).where(UnidadeCurricular.matriz_id == m.id, UnidadeCurricular.ativo.is_(True))
                                 .order_by(UnidadeCurricular.sequencia, UnidadeCurricular.codigo)).all()
                cod = {u.id: u.codigo for u in ucs}
                for u in ucs:
                    pre = [cod[p] for p in db.scalars(select(UCPrerequisito.prerequisito_id).where(UCPrerequisito.uc_id == u.id)) if p in cod]
                    res.append([c.codigo, c.nome, areas.get(c.area_id, ""), c.modalidade, c.tipo, c.carga_horaria_total, m.versao, u.codigo, u.nome,
                                u.carga_horaria_total, u.encontros_semanais, "; ".join(t.nome for t in u.tipos_ambiente), u.sequencia, "; ".join(pre),
                                u.certificacao_exigida])
    elif tipo == "ambientes":
        tipos = {t.id: t for t in db.scalars(select(TipoAmbiente))}
        q = select(Ambiente).where(Ambiente.ativo.is_(True)).order_by(Ambiente.codigo)
        for a in db.scalars(q):
            if unidades is not None and a.unidade_id not in unidades:
                continue
            t = tipos.get(a.tipo_ambiente_id)
            res.append([a.codigo, a.nome, unid.get(a.unidade_id), t.nome if t else "", "Sim" if t and t.especializado else "Não",
                        _areas_de(db, a.area_id, AmbienteArea, "ambiente_id", a.id, a.todas_areas, areas), a.capacidade_fisica, a.capacidade_operacional,
                        a.predio, "Sim" if a.acessivel else "Não", a.tempo_preparacao_min])
    elif tipo == "levantamento_instrutores":
        turnos = list(db.scalars(select(Turno).where(Turno.ativo.is_(True))))
        abrev = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
        for i in db.scalars(select(Instrutor).where(Instrutor.ativo.is_(True)).order_by(Instrutor.nome)):
            if unidades is not None and i.unidade_id not in unidades:
                continue
            disp = db.scalars(select(DisponibilidadeInstrutor).where(DisponibilidadeInstrutor.instrutor_id == i.id)).all()
            dias = sorted({d.dia_semana for d in disp})
            from app.services.regras_aula import jornada_dos_codigos
            cods = [t.codigo for t in turnos if t.unidade_id == i.unidade_id and any(d.hora_inicio <= t.hora_inicio and d.hora_fim >= t.hora_fim for d in disp)]
            jornada = jornada_dos_codigos(cods) or ""
            comp = db.scalars(select(CompetenciaInstrutor).where(CompetenciaInstrutor.instrutor_id == i.id, CompetenciaInstrutor.ativo.is_(True))).all()
            ucs = {u.id: u.codigo for u in db.scalars(select(UnidadeCurricular).where(UnidadeCurricular.id.in_({c.uc_id for c in comp})))} if comp else {}
            pleno = sorted({ucs[c.uc_id] for c in comp if c.uc_id in ucs and c.nivel != "supervisionado"})
            sup = sorted({ucs[c.uc_id] for c in comp if c.uc_id in ucs and c.nivel == "supervisionado"})
            cert = [h.nome for h in db.scalars(select(HabilitacaoInstrutor).where(HabilitacaoInstrutor.instrutor_id == i.id, HabilitacaoInstrutor.ativo.is_(True)))]
            res.append([i.matricula, i.nome, i.email, unid.get(i.unidade_id), _areas_de(db, i.area_id, InstrutorArea, "instrutor_id", i.id, i.todas_areas, areas),
                        (i.modalidade or "mensalista").capitalize(), "Básico" if i.nivel == "basico" else "Técnico", float(i.carga_horaria_contratual), float(i.limite_ch_letiva), float(i.max_horas_dia), jornada,
                        "; ".join(abrev[d] for d in dias), "; ".join(pleno), "; ".join(sup), "; ".join(cert), i.areas_competencia])
    elif tipo == "turmas":
        cursos = {c.id: c.codigo for c in db.scalars(select(Curso))}
        matrizes = {m.id: m.versao for m in db.scalars(select(MatrizCurricular))}
        abrev = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
        for t in db.scalars(select(Turma).where(Turma.ativo.is_(True), Turma.simulacao_id.is_(None)).order_by(Turma.codigo)):
            if unidades is not None and t.unidade_id not in unidades:
                continue
            res.append([t.codigo, t.nome, cursos.get(t.curso_id), matrizes.get(t.matriz_id), unid.get(t.unidade_id), turnos_nome.get(t.turno_id, ""),
                        t.qtd_alunos, t.data_inicio.strftime("%d/%m/%Y"), t.data_fim.strftime("%d/%m/%Y"),
                        "; ".join(abrev[d] for d in sorted(t.dias_possiveis or [])), areas.get(t.area_id, ""),
                        "; ".join(abrev[d] for d in sorted(t.dias_ead or []))])
    return res
