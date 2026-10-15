"""Importação de planilhas (XLSX/CSV) em duas etapas: prévia e confirmação.

1. O arquivo é lido, as colunas são mapeadas (sugestão automática por
   sinônimos, ajustável pelo usuário), as referências são resolvidas por
   identificadores confiáveis (códigos/matrículas) e cada linha recebe uma
   ação prevista: criar, atualizar, sem_alteracao ou erro.
2. Na confirmação, tudo é gravado em UMA transação. Aulas passam pela
   validação central; horários de cenários publicados nunca são alterados.

O hash SHA-256 do arquivo identifica reimportações: registros já existentes
são atualizados/ignorados pela chave natural, nunca duplicados.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from datetime import date, datetime, time

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import ConflitoValidacao, EntradaInvalida, NaoEncontrado
from app.models import (
    Alocacao,
    Ambiente,
    CenarioPlanejamento,
    CompetenciaInstrutor,
    Curso,
    DisponibilidadeInstrutor,
    Importacao,
    Instrutor,
    MatrizCurricular,
    PeriodoHorario,
    TipoAmbiente,
    Turma,
    TurmaUC,
    Turno,
    UCTipoAmbiente,
    Unidade,
    UnidadeCurricular,
)
from app.services import auditoria

DIAS = {"seg": 0, "segunda": 0, "ter": 1, "terca": 1, "qua": 2, "quarta": 2, "qui": 3, "quinta": 3, "sex": 4, "sexta": 4,
        "sab": 5, "sabado": 5, "dom": 6, "domingo": 6}


def norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


# --------------------------------------------------------------------------- conversões
def _str(v):
    if v is None:
        return None
    s = str(v).strip()
    if isinstance(v, float) and v.is_integer():
        s = str(int(v))
    return s or None


def _int(v):
    if v in (None, ""):
        return None
    return int(float(str(v).replace(",", ".")))


def _float(v):
    if v in (None, ""):
        return None
    return float(str(v).replace(",", "."))


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data inválida '{s}'")


def _time(v):
    if v in (None, ""):
        return None
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, float) and 0 <= v < 1:  # fração de dia do Excel
        m = round(v * 1440)
        return time(m // 60, m % 60)
    s = str(v).strip().lower().replace("h", ":")
    if s.endswith(":"):
        s += "00"
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass
    raise ValueError(f"horário inválido '{v}'")


def _bool(v):
    if v in (None, ""):
        return None
    return norm(v) in ("1", "sim", "s", "true", "verdadeiro", "x", "yes")


def _dia(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)) or str(v).strip().isdigit():
        n = int(float(v))
        return n - 1 if 1 <= n <= 7 else n  # aceita 1..7 (seg..dom)
    k = norm(v).replace("_feira", "").split("_")[0]
    if k[:3] in DIAS:
        return DIAS[k[:3]]
    raise ValueError(f"dia da semana inválido '{v}'")


def _lista_dias(v):
    """'seg;qua;sex', 'seg a sex', 'seg-sex', '1,2,3'."""
    if v in (None, ""):
        return None
    texto = norm(v).replace("_feira", "")
    m = re.fullmatch(r"([a-z]{3})[a-z]*_(?:a|ate)_([a-z]{3})[a-z]*", texto) or re.fullmatch(r"([a-z]{3})[a-z]*_([a-z]{3})[a-z]*", texto) \
        if re.search(r"_(a|ate)_|^[a-z]{3}[a-z]*_[a-z]{3}[a-z]*$", texto) and not re.search(r"[;,/]", str(v)) else None
    if m:
        d0, d1 = _dia(m.group(1)), _dia(m.group(2))
        return list(range(d0, d1 + 1))
    return [_dia(x) for x in re.split(r"[;,/ ]+", str(v)) if x.strip()]


CONV = {"str": _str, "int": _int, "float": _float, "date": _date, "time": _time, "bool": _bool, "dia": _dia, "dias": _lista_dias}

# --------------------------------------------------------------------------- especificações
# campo: (tipo, obrigatório, sinônimos)
SPECS: dict[str, dict] = {
    "unidades": {"chave": ["codigo"], "campos": {
        "codigo": ("str", True, ["cod", "codigo_unidade", "sigla"]), "nome": ("str", True, ["unidade", "nome_unidade"]),
        "endereco": ("str", False, []), "cidade": ("str", False, ["municipio"]), "responsavel": ("str", False, ["gestor"])}},
    "cursos": {"chave": ["codigo"], "campos": {
        "codigo": ("str", True, ["cod_curso", "codigo_curso"]), "nome": ("str", True, ["curso", "nome_curso"]),
        "modalidade": ("str", False, []), "tipo": ("str", False, ["tipo_curso"]), "carga_horaria_total": ("int", False, ["ch", "carga_horaria"]),
        "eixo_tecnologico": ("str", False, ["eixo"]), "unidade_codigo": ("str", False, ["unidade"]), "area": ("str", False, ["area_tecnologica"])}},
    "matrizes": {"chave": ["curso_codigo", "versao"], "campos": {
        "curso_codigo": ("str", True, ["curso"]), "versao": ("str", True, ["versao_matriz", "matriz"]),
        "vigencia_inicio": ("date", False, ["vigencia"])}},
    "ucs": {"chave": ["curso_codigo", "versao", "codigo"], "campos": {
        "curso_codigo": ("str", True, ["curso"]), "versao": ("str", True, ["versao_matriz", "matriz"]),
        "codigo": ("str", True, ["cod_uc", "codigo_uc"]), "nome": ("str", True, ["uc", "unidade_curricular", "nome_uc"]),
        "carga_horaria_total": ("int", True, ["ch", "carga_horaria", "ch_total"]), "ch_teorica": ("int", False, []), "ch_pratica": ("int", False, []),
        "duracao_bloco_periodos": ("int", False, ["periodos_por_encontro", "bloco"]), "encontros_semanais": ("int", False, ["encontros"]),
        "sequencia": ("int", False, ["ordem"]), "tipos_ambiente": ("str", False, ["tipo_ambiente", "ambiente_tipo"]),
        "certificacao_exigida": ("str", False, ["certificacao"]), "divide_grupos": ("bool", False, []), "capacidade_max": ("int", False, [])}},
    "turmas": {"chave": ["codigo"], "campos": {
        "codigo": ("str", True, ["turma", "cod_turma", "codigo_turma", "codigo_da_turma"]), "nome": ("str", False, ["nome_turma", "nome_da_turma"]),
        "curso_codigo": ("str", True, ["curso", "codigo_do_curso"]), "versao": ("str", True, ["versao_matriz", "matriz", "versao_da_matriz"]),
        "unidade_codigo": ("str", True, ["unidade"]), "turno_codigo": ("str", False, ["turno"]), "qtd_alunos": ("int", False, ["alunos", "quantidade_alunos"]),
        "data_inicio": ("date", True, ["inicio"]), "data_fim": ("date", True, ["fim", "termino"]), "dias_possiveis": ("dias", False, ["dias", "dias_de_aula"]),
        "num_grupos": ("int", False, ["grupos"]), "area": ("str", False, ["area_tecnologica"])}},
    "turma_ucs": {"chave": ["turma_codigo", "uc_codigo"], "campos": {
        "turma_codigo": ("str", True, ["turma"]), "uc_codigo": ("str", True, ["uc", "codigo_uc"]), "situacao": ("str", False, []),
        "data_inicio": ("date", False, ["inicio"]), "data_limite": ("date", False, ["limite", "fim"]), "ch_realizada": ("float", False, ["realizada"])}},
    "instrutores": {"chave": ["matricula"], "campos": {
        "matricula": ("str", True, ["registro", "identificacao", "id_institucional"]), "nome": ("str", True, ["instrutor", "docente", "professor"]),
        "email": ("str", False, ["e_mail"]), "unidade_codigo": ("str", True, ["unidade"]), "carga_horaria_contratual": ("float", False, ["ch_contratual", "jornada"]),
        "limite_ch_letiva": ("float", False, ["ch_letiva", "limite_letivo"]), "areas_competencia": ("str", False, ["areas"]),
        "area": ("str", False, ["area_tecnologica"])}},
    "competencias": {"chave": ["matricula", "uc_codigo"], "campos": {
        "matricula": ("str", True, ["instrutor"]), "uc_codigo": ("str", True, ["uc"]), "curso_codigo": ("str", False, ["curso"]),
        "versao": ("str", False, ["matriz"]), "nivel": ("str", False, []), "valido_ate": ("date", False, ["validade"])}},
    "ambientes": {"chave": ["codigo"], "campos": {
        "codigo": ("str", True, ["sala", "cod_ambiente", "codigo_da_sala"]), "nome": ("str", True, ["ambiente", "nome_ambiente", "nome_da_sala"]),
        "unidade_codigo": ("str", True, ["unidade"]),
        "tipo_codigo": ("str", True, ["tipo", "tipo_ambiente", "tipo_de_ambiente"]), "capacidade_fisica": ("int", True, ["capacidade", "lugares", "capacidade_maxima_lugares"]),
        "capacidade_operacional": ("int", False, ["capacidade_recomendada", "capacidade_recomendada_alunos"]), "predio": ("str", False, ["bloco", "predio_bloco"]),
        "acessivel": ("bool", False, ["acessibilidade", "acessivel_pcd"]), "especializado": ("bool", False, ["laboratorio_especializado"]),
        "tempo_preparacao_min": ("int", False, ["preparacao_entre_aulas_min"]),
        "area": ("str", False, ["area_tecnologica", "areas", "areas_que_usam"])}},
    # ---- planilhas únicas (cada linha grava também os registros ligados a ela)
    "estrutura_curricular": {"chave": ["curso_codigo", "versao", "uc_codigo"], "campos": {
        "curso_codigo": ("str", True, ["codigo_do_curso", "cod_curso"]), "curso_nome": ("str", True, ["nome_do_curso", "curso"]),
        "area": ("str", True, ["area_tecnologica"]), "modalidade": ("str", False, []), "tipo_curso": ("str", False, ["tipo_do_curso", "tipo_de_curso"]),
        "ch_curso": ("int", False, ["carga_horaria_do_curso_h", "carga_horaria_do_curso"]), "versao": ("str", False, ["versao_da_matriz", "matriz"]),
        "uc_codigo": ("str", True, ["codigo_da_uc", "codigo_da_unidade_curricular", "cod_uc"]),
        "uc_nome": ("str", True, ["nome_da_uc", "unidade_curricular", "nome_da_unidade_curricular"]),
        "carga_horaria_total": ("int", True, ["carga_horaria_da_uc_h", "carga_horaria_da_uc", "ch_da_uc"]),
        "encontros_semanais": ("int", False, ["dias_de_aula_por_semana", "encontros_por_semana"]),
        "tipos_ambiente": ("str", False, ["tipo_de_sala_exigido", "tipos_de_ambiente"]), "sequencia": ("int", False, ["ordem_na_matriz"]),
        "prerequisitos": ("str", False, ["pre_requisitos_codigos_das_ucs", "pre_requisitos"]),
        "certificacao_exigida": ("str", False, ["certificacao_exigida_do_instrutor"])}},
    "levantamento_instrutores": {"chave": ["matricula"], "campos": {
        "matricula": ("str", True, ["matricula_ou_registro"]), "nome": ("str", True, ["nome_completo", "instrutor"]), "email": ("str", False, ["e_mail"]),
        "unidade": ("str", True, ["unidade_principal"]), "area": ("str", True, ["areas_em_que_atua", "areas"]),
        "carga_horaria_contratual": ("float", False, ["carga_horaria_contratual_h_semana", "jornada"]),
        "limite_ch_letiva": ("float", False, ["limite_de_aulas_por_semana_h", "limite_letivo"]),
        "max_horas_dia": ("float", False, ["maximo_de_horas_por_dia"]),
        "turnos": ("str", False, ["turnos_de_trabalho", "turnos_disponiveis"]), "dias": ("dias", False, ["dias_disponiveis"]),
        "ucs_pleno": ("str", False, ["ucs_que_ministra_codigos", "ucs_que_ministra"]),
        "ucs_supervisionado": ("str", False, ["ucs_que_ministra_sob_supervisao_codigos", "ucs_sob_supervisao"]),
        "certificacoes": ("str", False, ["certificacoes_ex_nr_10"]), "areas_competencia": ("str", False, ["formacao_e_especialidades", "especialidades"])}},
    "disponibilidade_instrutores": {"chave": ["matricula", "dia_semana", "hora_inicio"], "campos": {
        "matricula": ("str", True, ["instrutor"]), "dia_semana": ("dia", True, ["dia"]), "hora_inicio": ("time", True, ["inicio"]),
        "hora_fim": ("time", True, ["fim"])}},
    "aulas": {"chave": ["turma_codigo", "uc_codigo", "data", "dia_semana", "hora_inicio", "grupo"], "campos": {
        "turma_codigo": ("str", True, ["turma"]), "uc_codigo": ("str", True, ["uc", "disciplina", "unidade_curricular"]),
        "matricula": ("str", False, ["instrutor_matricula"]), "instrutor_nome": ("str", False, ["instrutor", "docente", "professor"]),
        "ambiente_codigo": ("str", False, ["ambiente", "sala"]), "data": ("date", False, []), "dia_semana": ("dia", False, ["dia"]),
        "hora_inicio": ("time", True, ["inicio", "horario_inicio"]), "hora_fim": ("time", True, ["fim", "horario_fim", "termino"]),
        "grupo": ("int", False, []), "turno": ("str", False, [])}},
}


def ler_arquivo(nome: str, conteudo: bytes) -> dict[str, list[dict]]:
    """Retorna {nome_planilha: [linhas como dict]}. CSV vira uma planilha."""
    folhas = {}
    if nome.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
        for ws in wb.worksheets:
            linhas = list(ws.iter_rows(values_only=True))
            # cabeçalho = primeira linha com ao menos 2 células preenchidas
            idx = next((i for i, l in enumerate(linhas) if sum(1 for c in l if c not in (None, "")) >= 2), None)
            if idx is None:
                continue
            cab = [str(c).strip() if c is not None else f"col{j}" for j, c in enumerate(linhas[idx])]
            dados = [dict(zip(cab, l)) for l in linhas[idx + 1:] if any(c not in (None, "") for c in l)]
            folhas[ws.title] = dados
    elif nome.lower().endswith((".csv", ".txt")):
        texto = conteudo.decode("utf-8-sig", errors="replace")
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        folhas["csv"] = [dict(r) for r in csv.DictReader(io.StringIO(texto), dialect=dialeto)]
    else:
        raise EntradaInvalida("Formato não suportado. Use .xlsx ou .csv.")
    return folhas


def sugerir_mapeamento(tipo: str, cabecalhos: list[str]) -> dict[str, str | None]:
    """1ª passada: nome normalizado idêntico; 2ª: todos os termos do sinônimo
    contidos no cabeçalho (ex.: 'Nome do instrutor' → nome). O usuário revisa."""
    campos = SPECS[tipo]["campos"]
    mapa = {c: None for c in campos}
    usados = set()
    ncab = {norm(c): c for c in cabecalhos}
    for campo, (_t, _o, sin) in campos.items():
        for cand in [campo, *sin]:
            if norm(cand) in ncab and ncab[norm(cand)] not in usados:
                mapa[campo] = ncab[norm(cand)]
                usados.add(mapa[campo])
                break
    for campo, (_t, _o, sin) in campos.items():
        if mapa[campo]:
            continue
        for cand in [campo, *sin]:
            termos = set(norm(cand).split("_"))
            achado = next((orig for n, orig in ncab.items() if orig not in usados and termos <= set(n.split("_"))), None)
            if achado:
                mapa[campo] = achado
                usados.add(achado)
                break
    return mapa


def detectar_tipo(cabecalhos: list[str]) -> str:
    melhor, pts = None, -1
    for tipo in SPECS:
        m = sugerir_mapeamento(tipo, cabecalhos)
        obrig = [c for c, (_t, o, _s) in SPECS[tipo]["campos"].items() if o]
        if any(m[c] is None for c in obrig):
            continue
        p = sum(1 for v in m.values() if v)
        if p > pts:
            melhor, pts = tipo, p
    if melhor is None:
        raise EntradaInvalida("Não foi possível identificar o tipo da planilha pelos cabeçalhos; selecione o tipo e ajuste o mapeamento.")
    return melhor


# --------------------------------------------------------------------------- resolução
class Resolvedor:
    def __init__(self, db: Session):
        self.db = db
        self.c = {}

    def _get(self, k, fn):
        if k not in self.c:
            self.c[k] = fn()
        return self.c[k]

    def unidade(self, cod):
        def busca():
            alvo = norm(cod)
            return next((u.id for u in self.db.scalars(select(Unidade)) if alvo in (norm(u.codigo), norm(u.nome), norm(u.nome.replace("Unidade ", "")))), None)
        return self._get(("u", cod), busca)

    def ucs_por_codigo(self, cod):
        return self._get(("ucc", cod), lambda: self.db.scalars(select(UnidadeCurricular.id).where(func.lower(UnidadeCurricular.codigo) == cod.lower(),
                                                                                               UnidadeCurricular.ativo.is_(True))).all())

    def area(self, nome_ou_cod):
        def busca():
            from app.models import Area
            alvo = norm(nome_ou_cod)
            return next(((a.id, a.unidade_id) for a in self.db.scalars(select(Area)) if alvo in (norm(a.codigo), norm(a.nome))), None)
        return self._get(("a", nome_ou_cod), busca)

    def curso(self, cod):
        return self._get(("c", cod), lambda: self.db.scalar(select(Curso.id).where(func.lower(Curso.codigo) == cod.lower())))

    def matriz(self, curso_cod, versao):
        cid = self.curso(curso_cod) if curso_cod else None
        return self._get(("m", cid, versao), lambda: self.db.scalar(select(MatrizCurricular.id).where(MatrizCurricular.curso_id == cid, MatrizCurricular.versao == versao)) if cid else None)

    def turno(self, unidade_id, cod):
        return self._get(("t", unidade_id, cod), lambda: self.db.scalar(select(Turno.id).where(Turno.unidade_id == unidade_id, (func.lower(Turno.codigo) == cod.lower()) | (func.lower(Turno.nome) == cod.lower()))))

    def tipo_amb(self, cod):
        return self._get(("ta", cod), lambda: self.db.scalar(select(TipoAmbiente.id).where((func.lower(TipoAmbiente.codigo) == cod.lower()) | (func.lower(TipoAmbiente.nome) == cod.lower()))))

    def turma(self, cod):
        return self._get(("tu", cod), lambda: self.db.scalars(select(Turma).where(func.lower(Turma.codigo) == cod.lower(), Turma.simulacao_id.is_(None))).first())

    def uc_da_turma(self, turma: Turma, cod):
        return self._get(("uct", turma.id, cod), lambda: self.db.scalar(select(UnidadeCurricular.id).where(UnidadeCurricular.matriz_id == turma.matriz_id, func.lower(UnidadeCurricular.codigo) == cod.lower())))

    def instrutor(self, matricula=None, nome=None):
        if matricula:
            return self._get(("i", matricula), lambda: self.db.scalar(select(Instrutor.id).where(Instrutor.matricula == matricula))), None
        if nome:
            ids = self.db.scalars(select(Instrutor.id).where(func.lower(Instrutor.nome) == nome.strip().lower())).all()
            if len(ids) == 1:
                return ids[0], f"Instrutor identificado pelo nome '{nome}' (sem matrícula): revise a correspondência."
            if len(ids) > 1:
                return None, f"Nome '{nome}' corresponde a {len(ids)} instrutores: informe a matrícula."
        return None, None

    def ambiente(self, cod):
        return self._get(("a", cod), lambda: self.db.scalar(select(Ambiente.id).where(func.lower(Ambiente.codigo) == cod.lower())))


def _converter(tipo: str, linha: dict, mapa: dict) -> tuple[dict, list[str]]:
    vals, erros = {}, []
    for campo, (t, obrig, _s) in SPECS[tipo]["campos"].items():
        col = mapa.get(campo)
        bruto = linha.get(col) if col else None
        try:
            v = CONV[t](bruto)
        except (ValueError, TypeError) as e:
            erros.append(f"{campo}: {e}")
            v = None
        if obrig and v in (None, "") and not any(e.startswith(campo) for e in erros):
            erros.append(f"{campo}: obrigatório")
        vals[campo] = v
    return vals, erros


def _area(r, v: dict, d: dict, E: list, A: list, obrigatoria: bool = True, multiplas: bool = False):
    """Coluna 'área': nome ou código de uma das áreas fixas, da mesma unidade do registro.
    Para instrutores e ambientes aceita várias (separadas por ';' ou ',') — a primeira é a principal —
    ou 'TODAS' (atende todas as áreas)."""
    texto = (v.get("area") or "").strip()
    if not texto:
        if obrigatoria:
            A.append("área não informada: defina-a depois em Cadastros (coluna 'area', ex.: AUTOMAÇÃO, LOGÍSTICA)")
        return
    if multiplas and norm(texto) in ("todas", "todas_as_areas", "todas as areas"):
        d["todas_areas"] = True
        return
    nomes = [x.strip() for x in re.split(r"[;,|]+", texto) if x.strip()] if multiplas else [texto]
    ids = []
    for nome in nomes:
        achada = r.area(nome)
        if not achada:
            E.append(f"área '{nome}' não existe (use uma das áreas fixas da unidade)")
            return
        area_id, unidade_area = achada
        if d.get("unidade_id") and d["unidade_id"] != unidade_area:
            E.append(f"a área '{nome}' não pertence à unidade informada")
            return
        d.setdefault("unidade_id", unidade_area)
        ids.append(area_id)
    d["area_id"] = ids[0]
    if multiplas:
        d["_areas"] = ids[1:]


def _lista(v) -> list[str]:
    return [x.strip() for x in re.split(r"[;,|\n]+", str(v or "")) if x.strip()]


def _preparar_estrutura(db, r, v, E, A):
    """Uma linha = uma UC; o curso e a matriz são criados/atualizados junto."""
    versao = v["versao"] or "1"
    curso = {"codigo": v["curso_codigo"], "nome": v["curso_nome"]}
    for k_src, k_dst in (("modalidade", "modalidade"), ("tipo_curso", "tipo"), ("ch_curso", "carga_horaria_total")):
        if v.get(k_src) is not None:
            curso[k_dst] = v[k_src]
    _area(r, v, curso, E, A)
    mid = r.matriz(v["curso_codigo"], versao)
    ex = db.scalars(select(UnidadeCurricular).where(UnidadeCurricular.matriz_id == mid, func.lower(UnidadeCurricular.codigo) == v["uc_codigo"].lower())).first() if mid else None
    d = {"codigo": v["uc_codigo"], "nome": v["uc_nome"], "carga_horaria_total": v["carga_horaria_total"],
         "encontros_semanais": v["encontros_semanais"] or 1, "duracao_bloco_periodos": 4}
    for k in ("sequencia", "certificacao_exigida"):
        if v.get(k) is not None:
            d[k] = v[k]
    if ex is not None and mid and db.scalar(select(func.count()).select_from(Turma).where(Turma.matriz_id == mid)):
        if any(getattr(ex, k) != d[k] for k in ("carga_horaria_total",) if k in d):
            E.append("a matriz já é usada por turmas: para mudar a carga horária, use uma nova versão da matriz (coluna 'Versão da matriz')")
    if ex is not None:
        d["matriz_id"] = ex.matriz_id
    d["_curso"], d["_versao"] = curso, versao
    d["_tipos_nomes"] = _lista(v["tipos_ambiente"])
    for t in d["_tipos_nomes"]:
        if not r.tipo_amb(t):
            A.append(f"tipo de sala '{t}' ainda não existe: será criado")
    d["_prereq"] = _lista(v["prerequisitos"])
    return d, ex, E, A


def _preparar_levantamento(db, r, v, E, A):
    """Uma linha = um instrutor com áreas, disponibilidade (turnos × dias), competências e certificações."""
    uid = r.unidade(v["unidade"])
    if not uid:
        E.append(f"unidade '{v['unidade']}' não encontrada (use Beira Mar ou Porto)")
    d = {"matricula": v["matricula"], "nome": v["nome"], "unidade_id": uid}
    for k in ("email", "carga_horaria_contratual", "limite_ch_letiva", "max_horas_dia", "areas_competencia"):
        if v.get(k) is not None:
            d[k] = v[k]
    if d.get("limite_ch_letiva") and d.get("carga_horaria_contratual") and d["limite_ch_letiva"] > d["carga_horaria_contratual"]:
        E.append("limite de aulas por semana maior que a carga contratual")
    _area(r, v, d, E, A, multiplas=True)
    disp = []
    from app.services.regras_aula import JORNADAS, horas_dia_pela_carga, normalizar_jornada, validar_jornada
    if d.get("max_horas_dia") is None and d.get("carga_horaria_contratual"):
        d["max_horas_dia"] = horas_dia_pela_carga(d["carga_horaria_contratual"])
    if v["turnos"]:
        # turnos de trabalho: M, V ou N (um turno) ou MV / VN (40 h = dois turnos); aceita também "Manhã; Tarde"
        j = normalizar_jornada(str(v["turnos"]))
        erro = "turnos de trabalho inválidos: use M, V, N, MV ou VN" if j is None else validar_jornada(j, d.get("max_horas_dia") or 8)
        if erro:
            E.append(erro[0].lower() + erro[1:] if erro[0].isupper() else erro)
        else:
            dias = v["dias"] if v["dias"] is not None else [0, 1, 2, 3, 4]
            for cod in JORNADAS[j]:
                tid = r.turno(uid, {"M": "Manhã", "T": "Tarde", "N": "Noite"}[cod]) if uid else None
                turno = db.get(Turno, tid) if tid else None
                if not turno:
                    E.append("a unidade não tem os turnos desta jornada cadastrados")
                    continue
                disp += [[dd, turno.hora_inicio.strftime("%H:%M"), turno.hora_fim.strftime("%H:%M")] for dd in dias]
    else:
        A.append("turnos de trabalho não informados: cadastre depois (M, V, N, MV ou VN)")
    comps = []
    for col, nivel in (("ucs_pleno", "pleno"), ("ucs_supervisionado", "supervisionado")):
        for cod in _lista(v[col]):
            ids = r.ucs_por_codigo(cod)
            if not ids:
                E.append(f"UC '{cod}' não cadastrada: importe antes a estrutura curricular")
            comps += [[i, nivel] for i in ids]
    if not comps:
        A.append("nenhuma UC informada: o instrutor não será sugerido para aulas até ter competências")
    d["_disp"], d["_comp"], d["_cert"] = disp, comps, _lista(v["certificacoes"])
    ex = db.scalars(select(Instrutor).where(Instrutor.matricula == v["matricula"])).first()
    return d, ex, E, A


def _preparar(db: Session, r: Resolvedor, tipo: str, v: dict) -> tuple[dict | None, object | None, list[str], list[str]]:
    """Retorna (dados para o modelo, registro existente, erros, avisos)."""
    E, A = [], []
    if tipo == "unidades":
        ex = db.scalars(select(Unidade).where(func.lower(Unidade.codigo) == v["codigo"].lower())).first()
        if ex is None:
            E.append("as unidades são fixas da instituição (BEIRA-MAR e PORTO); não é possível incluir outra")
        # só dados complementares; código e nome são fixos
        return {k: v[k] for k in ("endereco", "cidade", "responsavel") if v[k] is not None}, ex, E, A
    if tipo == "cursos":
        d = {k: v[k] for k in ("codigo", "nome", "modalidade", "tipo", "carga_horaria_total", "eixo_tecnologico") if v[k] is not None}
        if v["unidade_codigo"]:
            d["unidade_id"] = r.unidade(v["unidade_codigo"]) or E.append(f"unidade '{v['unidade_codigo']}' não cadastrada")
        _area(r, v, d, E, A)
        return d, db.scalars(select(Curso).where(func.lower(Curso.codigo) == v["codigo"].lower())).first(), E, A
    if tipo == "matrizes":
        cid = r.curso(v["curso_codigo"])
        if not cid:
            return None, None, [f"curso '{v['curso_codigo']}' não cadastrado"], A
        ex = db.scalars(select(MatrizCurricular).where(MatrizCurricular.curso_id == cid, MatrizCurricular.versao == v["versao"])).first()
        return {"curso_id": cid, "versao": v["versao"], "vigencia_inicio": v["vigencia_inicio"]}, ex, E, A
    if tipo == "ucs":
        mid = r.matriz(v["curso_codigo"], v["versao"])
        if not mid:
            return None, None, [f"matriz {v['curso_codigo']} v{v['versao']} não cadastrada"], A
        ex = db.scalars(select(UnidadeCurricular).where(UnidadeCurricular.matriz_id == mid, func.lower(UnidadeCurricular.codigo) == v["codigo"].lower())).first()
        if ex and db.scalar(select(func.count()).select_from(Turma).where(Turma.matriz_id == mid)):
            E.append("a matriz já está em uso por turmas: crie uma nova versão da matriz para alterar UCs")
        d = {k: v[k] for k in ("codigo", "nome", "carga_horaria_total", "ch_teorica", "ch_pratica", "duracao_bloco_periodos", "encontros_semanais",
                               "sequencia", "certificacao_exigida", "divide_grupos", "capacidade_max") if v[k] is not None}
        d["matriz_id"] = mid
        tipos = []
        for t in re.split(r"[;,]+", v["tipos_ambiente"] or ""):
            if t.strip():
                tid = r.tipo_amb(t.strip())
                tipos.append(tid) if tid else E.append(f"tipo de ambiente '{t.strip()}' não cadastrado")
        d["_tipos"] = tipos
        return d, ex, E, A
    if tipo == "turmas":
        mid = r.matriz(v["curso_codigo"], v["versao"])
        uid = r.unidade(v["unidade_codigo"])
        if not mid:
            E.append(f"matriz {v['curso_codigo']} v{v['versao']} não cadastrada")
        if not uid:
            E.append(f"unidade '{v['unidade_codigo']}' não cadastrada")
        d = {"codigo": v["codigo"], "nome": v["nome"] or v["codigo"], "curso_id": r.curso(v["curso_codigo"]), "matriz_id": mid, "unidade_id": uid,
             "data_inicio": v["data_inicio"], "data_fim": v["data_fim"]}
        for k in ("qtd_alunos", "dias_possiveis", "num_grupos"):
            if v[k] is not None:
                d[k] = v[k]
        if v["turno_codigo"] and uid:
            d["turno_id"] = r.turno(uid, v["turno_codigo"]) or E.append(f"turno '{v['turno_codigo']}' não cadastrado na unidade")
        if v["data_inicio"] and v["data_fim"] and v["data_fim"] < v["data_inicio"]:
            E.append("data_fim anterior a data_inicio")
        if not v["area"] and d.get("curso_id"):
            curso = db.get(Curso, d["curso_id"])
            if curso and curso.area_id:
                d["area_id"] = curso.area_id  # herda a área do curso
        _area(r, v, d, E, A, obrigatoria=not d.get("area_id"))
        return d, r.turma(v["codigo"]), E, A
    if tipo == "turma_ucs":
        t = r.turma(v["turma_codigo"])
        if not t:
            return None, None, [f"turma '{v['turma_codigo']}' não cadastrada"], A
        uid = r.uc_da_turma(t, v["uc_codigo"])
        if not uid:
            return None, None, [f"UC '{v['uc_codigo']}' não pertence à matriz da turma"], A
        ex = db.scalars(select(TurmaUC).where(TurmaUC.turma_id == t.id, TurmaUC.uc_id == uid)).first()
        d = {"turma_id": t.id, "uc_id": uid}
        for k in ("situacao", "data_inicio", "data_limite", "ch_realizada"):
            if v[k] is not None:
                d[k] = v[k]
        return d, ex, E, A
    if tipo == "instrutores":
        uid = r.unidade(v["unidade_codigo"]) or E.append(f"unidade '{v['unidade_codigo']}' não cadastrada")
        d = {k: v[k] for k in ("matricula", "nome", "email", "carga_horaria_contratual", "limite_ch_letiva", "areas_competencia") if v[k] is not None}
        d["unidade_id"] = uid
        _area(r, v, d, E, A, multiplas=True)
        return d, db.scalars(select(Instrutor).where(Instrutor.matricula == v["matricula"])).first(), E, A
    if tipo == "competencias":
        iid, aviso = r.instrutor(v["matricula"])
        if not iid:
            E.append(f"instrutor '{v['matricula']}' não cadastrado")
        q = select(UnidadeCurricular.id).where(func.lower(UnidadeCurricular.codigo) == v["uc_codigo"].lower())
        if v["curso_codigo"] and v["versao"]:
            q = q.where(UnidadeCurricular.matriz_id == r.matriz(v["curso_codigo"], v["versao"]))
        ucs = db.scalars(q).all()
        if len(ucs) != 1:
            E.append(f"UC '{v['uc_codigo']}' {'não encontrada' if not ucs else 'ambígua (informe curso e versão da matriz)'}")
            return None, None, E, A
        nivel = norm(v["nivel"] or "pleno")
        if nivel not in ("pleno", "supervisionado", "em_formacao"):
            E.append(f"nível '{v['nivel']}' inválido (pleno, supervisionado, em_formacao)")
        ex = db.scalars(select(CompetenciaInstrutor).where(CompetenciaInstrutor.instrutor_id == iid, CompetenciaInstrutor.uc_id == ucs[0])).first() if iid else None
        return {"instrutor_id": iid, "uc_id": ucs[0], "nivel": nivel, "valido_ate": v["valido_ate"]}, ex, E, A
    if tipo == "estrutura_curricular":
        return _preparar_estrutura(db, r, v, E, A)
    if tipo == "levantamento_instrutores":
        return _preparar_levantamento(db, r, v, E, A)
    if tipo == "ambientes":
        uid = r.unidade(v["unidade_codigo"]) or E.append(f"unidade '{v['unidade_codigo']}' não cadastrada")
        tid = r.tipo_amb(v["tipo_codigo"])
        d = {"codigo": v["codigo"], "nome": v["nome"], "unidade_id": uid, "tipo_ambiente_id": tid, "capacidade_fisica": v["capacidade_fisica"],
             "capacidade_operacional": v["capacidade_operacional"] or v["capacidade_fisica"]}
        for k in ("predio", "acessivel", "tempo_preparacao_min"):
            if v.get(k) is not None:
                d[k] = v[k]
        if not tid:
            A.append(f"tipo de ambiente '{v['tipo_codigo']}' ainda não existe: será criado")
            d["_novo_tipo"] = {"nome": v["tipo_codigo"], "especializado": bool(v.get("especializado"))}
        if d["capacidade_operacional"] and d["capacidade_fisica"] and d["capacidade_operacional"] > d["capacidade_fisica"]:
            E.append("capacidade operacional maior que a física")
        _area(r, v, d, E, A, multiplas=True)
        return d, db.scalars(select(Ambiente).where(func.lower(Ambiente.codigo) == v["codigo"].lower())).first(), E, A
    if tipo == "disponibilidade_instrutores":
        iid, _ = r.instrutor(v["matricula"])
        if not iid:
            return None, None, [f"instrutor '{v['matricula']}' não cadastrado"], A
        if v["hora_fim"] and v["hora_inicio"] and v["hora_fim"] <= v["hora_inicio"]:
            E.append("hora_fim deve ser posterior a hora_inicio")
        ex = db.scalars(select(DisponibilidadeInstrutor).where(DisponibilidadeInstrutor.instrutor_id == iid,
                                                               DisponibilidadeInstrutor.dia_semana == v["dia_semana"],
                                                               DisponibilidadeInstrutor.hora_inicio == v["hora_inicio"])).first()
        return {"instrutor_id": iid, "dia_semana": v["dia_semana"], "hora_inicio": v["hora_inicio"], "hora_fim": v["hora_fim"]}, ex, E, A
    if tipo == "aulas":
        t = r.turma(v["turma_codigo"])
        if not t:
            return None, None, [f"turma '{v['turma_codigo']}' não cadastrada"], A
        ucid = r.uc_da_turma(t, v["uc_codigo"])
        tuc = db.scalars(select(TurmaUC).where(TurmaUC.turma_id == t.id, TurmaUC.uc_id == ucid)).first() if ucid else None
        if not tuc:
            E.append(f"UC '{v['uc_codigo']}' não está vinculada à turma {t.codigo}")
        iid, aviso = r.instrutor(v["matricula"], v["instrutor_nome"])
        if aviso:
            (A if iid else E).append(aviso)
        if (v["matricula"] or v["instrutor_nome"]) and not iid and not aviso:
            E.append(f"instrutor '{v['matricula'] or v['instrutor_nome']}' não cadastrado")
        aid = r.ambiente(v["ambiente_codigo"]) if v["ambiente_codigo"] else None
        if v["ambiente_codigo"] and not aid:
            E.append(f"ambiente '{v['ambiente_codigo']}' não cadastrado")
        if v["data"] is None and v["dia_semana"] is None:
            E.append("informe a data ou o dia da semana")
        if v["hora_inicio"] and v["hora_fim"] and v["hora_fim"] <= v["hora_inicio"]:
            E.append("hora_fim deve ser posterior a hora_inicio")
        return {"turma_uc_id": tuc.id if tuc else None, "turma_id": t.id, "unidade_id": t.unidade_id, "instrutor_id": iid, "ambiente_id": aid,
                "data": v["data"], "dia_semana": v["dia_semana"] if v["data"] is None else v["data"].weekday(),
                "hora_inicio": v["hora_inicio"], "hora_fim": v["hora_fim"], "grupo": v["grupo"]}, None, E, A
    raise EntradaInvalida(f"Tipo de importação desconhecido: {tipo}")


def _json(d):
    return auditoria._json({k: v for k, v in (d or {}).items() if not k.startswith("_")})


def gerar_previa(db: Session, ctx, nome: str, conteudo: bytes, tipo: str | None = None, planilha: str | None = None,
                 mapeamento: dict | None = None, cenario_id: int | None = None, importacao_id: int | None = None) -> Importacao:
    h = hashlib.sha256(conteudo).hexdigest()
    folhas = ler_arquivo(nome, conteudo)
    if not folhas:
        raise EntradaInvalida("Nenhuma planilha com dados foi encontrada.")
    planilha = planilha if planilha in folhas else next(iter(folhas))
    linhas = folhas[planilha]
    cab = list(linhas[0].keys()) if linhas else []
    tipo = tipo or detectar_tipo(cab)
    if tipo not in SPECS:
        raise EntradaInvalida(f"Tipo '{tipo}' inválido. Opções: {', '.join(SPECS)}")
    mapa = {**sugerir_mapeamento(tipo, cab), **(mapeamento or {})}
    r = Resolvedor(db)
    itens, vistos = [], {}
    for n, linha in enumerate(linhas, start=2):
        if any(str(x or "").strip().upper().startswith("EXEMPLO") for x in list(linha.values())[:2]):
            continue  # linhas de exemplo dos modelos
        vals, erros = _converter(tipo, linha, mapa)
        avisos = []
        dados, existente = None, None
        if not erros:
            dados, existente, e2, avisos = _preparar(db, r, tipo, vals)
            erros += e2
        chave = tuple(str(vals.get(k)).lower() for k in SPECS[tipo]["chave"])
        if chave in vistos:
            erros.append(f"registro duplicado no arquivo (mesma chave da linha {vistos[chave]})")
        else:
            vistos[chave] = n
        if erros:
            acao = "erro"
        elif tipo == "aulas":
            acao = "criar"
        elif existente is None:
            acao = "criar"
        else:
            atual = auditoria.snapshot(existente)
            mud = {k: v for k, v in _json(dados).items() if atual.get(k) != v}
            acao = "atualizar" if mud else "sem_alteracao"
        itens.append({"linha": n, "acao": acao, "erros": erros, "avisos": avisos, "valores": _json(vals), "dados": _json(dados),
                      "tipos": (dados or {}).get("_tipos"), "areas_extra": (dados or {}).get("_areas"),
                      "extras": {k: v for k, v in (dados or {}).items() if k.startswith("_") and k not in ("_tipos", "_areas")}, "existente_id": getattr(existente, "id", None) if existente is not None else None})
    anterior = db.scalars(select(Importacao).where(Importacao.arquivo_hash == h, Importacao.status == "confirmada")).first()
    resumo = {a: sum(1 for i in itens if i["acao"] == a) for a in ("criar", "atualizar", "sem_alteracao", "erro")}
    previa = {"planilhas": list(folhas), "planilha": planilha, "cabecalhos": cab, "campos": {k: {"tipo": t, "obrigatorio": o} for k, (t, o, _s) in SPECS[tipo]["campos"].items()},
              "itens": itens, "resumo": resumo,
              "arquivo_ja_processado": {"importacao_id": anterior.id, "em": anterior.confirmado_em.isoformat() if anterior.confirmado_em else None} if anterior else None}
    imp = db.get(Importacao, importacao_id) if importacao_id else None
    if imp is None:
        imp = Importacao(arquivo_nome=nome, arquivo_hash=h, criado_por_id=getattr(ctx, "id", None))
        db.add(imp)
    imp.tipo, imp.mapeamento, imp.previa, imp.status, imp.cenario_id = tipo, mapa, previa, "previa", cenario_id
    imp.conteudo = conteudo
    db.flush()
    return imp


def confirmar(db: Session, ctx, importacao_id: int, ignorar_erros: bool = False, correcoes: dict | None = None) -> dict:
    """Grava a importação em uma transação. Com `ignorar_erros`, linhas com
    erro são descartadas (e relatadas); sem ele, qualquer erro impede a gravação."""
    imp = db.scalars(select(Importacao).where(Importacao.id == importacao_id).with_for_update()).first()
    if imp is None:
        raise NaoEncontrado("Importação não encontrada.")
    if imp.status != "previa":
        raise ConflitoValidacao(f"Importação já está '{imp.status}'.")
    itens = imp.previa["itens"]
    erros = [i for i in itens if i["acao"] == "erro"]
    if erros and not ignorar_erros:
        raise ConflitoValidacao(f"{len(erros)} linha(s) com erro. Corrija o arquivo/mapeamento ou confirme ignorando as linhas com erro.",
                                detalhes={"erros": erros[:100]})
    modelos = {"unidades": Unidade, "cursos": Curso, "matrizes": MatrizCurricular, "ucs": UnidadeCurricular, "turmas": Turma, "turma_ucs": TurmaUC,
               "instrutores": Instrutor, "levantamento_instrutores": Instrutor, "competencias": CompetenciaInstrutor, "ambientes": Ambiente, "disponibilidade_instrutores": DisponibilidadeInstrutor}
    rel = {"criados": 0, "atualizados": 0, "sem_alteracao": 0, "ignorados_por_erro": len(erros), "aulas_bloqueadas": [], "avisos": []}
    if imp.tipo == "aulas":
        rel.update(_confirmar_aulas(db, ctx, imp, [i for i in itens if i["acao"] != "erro"]))
    elif imp.tipo == "estrutura_curricular":
        rel.update(_confirmar_estrutura(db, ctx, imp, itens))
    else:
        Modelo = modelos[imp.tipo]
        conv = {c.key: c for c in Modelo.__table__.columns}
        for it in itens:
            if it["acao"] in ("erro", "sem_alteracao"):
                rel["sem_alteracao"] += it["acao"] == "sem_alteracao"
                continue
            dados = {}
            for k, v in it["dados"].items():
                col = conv.get(k)
                if col is not None and v is not None:
                    pt = col.type.python_type if hasattr(col.type, "python_type") else None
                    try:
                        if pt is date and isinstance(v, str):
                            v = date.fromisoformat(v)
                        elif pt is time and isinstance(v, str):
                            v = time.fromisoformat(v)
                    except NotImplementedError:
                        pass
                dados[k] = v
            ext = it.get("extras") or {}
            if imp.tipo == "ambientes" and ext.get("_novo_tipo"):
                dados["tipo_ambiente_id"] = _tipo_ambiente(db, ext["_novo_tipo"]["nome"], ext["_novo_tipo"]["especializado"])
            if it["acao"] == "criar":
                obj = Modelo(**dados)
                db.add(obj)
                db.flush()
                rel["criados"] += 1
                auditoria.registrar(db, ctx, "importacao.criar", Modelo.__tablename__, getattr(obj, "id", None), depois=dados, origem=f"importacao:{imp.id}")
            else:
                obj = db.get(Modelo, it["existente_id"])
                antes = auditoria.snapshot(obj)
                for k, v in dados.items():
                    setattr(obj, k, v)
                rel["atualizados"] += 1
                auditoria.registrar(db, ctx, "importacao.atualizar", Modelo.__tablename__, getattr(obj, "id", None), antes=antes, depois=dados, origem=f"importacao:{imp.id}")
            if imp.tipo == "ucs" and it.get("tipos") is not None:
                db.flush()
                for tid in it["tipos"]:
                    if not db.get(UCTipoAmbiente, (obj.id, tid)):
                        db.add(UCTipoAmbiente(uc_id=obj.id, tipo_ambiente_id=tid))
            if imp.tipo == "levantamento_instrutores":
                _gravar_levantamento(db, obj, ext)
            if imp.tipo == "turmas" and it["acao"] == "criar":
                # mesmas regras do cadastro pela tela: as UCs da matriz viram ofertas da turma
                for uc in db.scalars(select(UnidadeCurricular).where(UnidadeCurricular.matriz_id == obj.matriz_id, UnidadeCurricular.ativo.is_(True))):
                    db.add(TurmaUC(turma_id=obj.id, uc_id=uc.id, situacao="prevista"))
            if imp.tipo in ("instrutores", "ambientes", "levantamento_instrutores") and it.get("areas_extra"):
                from app.models import AmbienteArea, InstrutorArea
                db.flush()
                M, chave = (AmbienteArea, "ambiente_id") if imp.tipo == "ambientes" else (InstrutorArea, "instrutor_id")
                for aid in it["areas_extra"]:
                    if aid != obj.area_id and not db.get(M, (obj.id, aid)):
                        db.add(M(**{chave: obj.id, "area_id": aid}))
            rel["avisos"] += [f"linha {it['linha']}: {a}" for a in it["avisos"]]
    imp.status = "confirmada"
    imp.confirmado_em = datetime.now()
    imp.conteudo = None  # o arquivo não é mantido após a confirmação (apenas hash, prévia e relatório)
    imp.relatorio = rel
    auditoria.registrar(db, ctx, "importacao.confirmar", "importacoes", imp.id, depois={"arquivo": imp.arquivo_nome, "hash": imp.arquivo_hash, **{k: v for k, v in rel.items() if k != "aulas_bloqueadas"}})
    return rel


def _tipo_ambiente(db, nome: str, especializado: bool | None = None) -> int:
    t = db.scalars(select(TipoAmbiente).where((func.lower(TipoAmbiente.nome) == nome.lower()) | (func.lower(TipoAmbiente.codigo) == nome.lower()))).first()
    if t is None:
        esp = especializado if especializado is not None else ("lab" in norm(nome) or "oficina" in norm(nome))
        t = TipoAmbiente(codigo=norm(nome).upper()[:30], nome=nome, especializado=esp)
        db.add(t)
        db.flush()
    return t.id


def _confirmar_estrutura(db: Session, ctx, imp: Importacao, itens: list[dict]) -> dict:
    """Cria/atualiza cursos, matrizes e UCs (com tipos de sala e pré-requisitos) a partir da planilha única."""
    rel = {"criados": 0, "atualizados": 0, "sem_alteracao": 0, "cursos": 0, "matrizes": 0}
    por_matriz: dict[int, dict[str, int]] = {}
    prereqs = []
    for it in itens:
        if it["acao"] == "erro":
            continue
        ext = it["extras"]
        c = ext["_curso"]
        curso = db.scalars(select(Curso).where(func.lower(Curso.codigo) == c["codigo"].lower())).first()
        if curso is None:
            curso = Curso(situacao="ativo", ativo=True, **c)
            db.add(curso)
            db.flush()
            rel["cursos"] += 1
            auditoria.registrar(db, ctx, "importacao.criar", "cursos", curso.id, depois=c, origem=f"importacao:{imp.id}")
        else:
            for k, v in c.items():
                if v is not None and getattr(curso, k) != v:
                    setattr(curso, k, v)
        mat = db.scalars(select(MatrizCurricular).where(MatrizCurricular.curso_id == curso.id, MatrizCurricular.versao == ext["_versao"])).first()
        if mat is None:
            mat = MatrizCurricular(curso_id=curso.id, versao=ext["_versao"], situacao="vigente", ativo=True)
            db.add(mat)
            db.flush()
            rel["matrizes"] += 1
        dados = dict(it["dados"])
        dados["matriz_id"] = mat.id
        uc = db.scalars(select(UnidadeCurricular).where(UnidadeCurricular.matriz_id == mat.id, func.lower(UnidadeCurricular.codigo) == dados["codigo"].lower())).first()
        if uc is None:
            uc = UnidadeCurricular(**dados)
            db.add(uc)
            db.flush()
            rel["criados"] += 1
            auditoria.registrar(db, ctx, "importacao.criar", "unidades_curriculares", uc.id, depois=dados, origem=f"importacao:{imp.id}")
        elif it["acao"] == "sem_alteracao":
            rel["sem_alteracao"] += 1
        else:
            antes = auditoria.snapshot(uc)
            for k, v in dados.items():
                setattr(uc, k, v)
            rel["atualizados"] += 1
            auditoria.registrar(db, ctx, "importacao.atualizar", "unidades_curriculares", uc.id, antes=antes, depois=dados, origem=f"importacao:{imp.id}")
        for nome in ext.get("_tipos_nomes") or []:
            tid = _tipo_ambiente(db, nome)
            if not db.get(UCTipoAmbiente, (uc.id, tid)):
                db.add(UCTipoAmbiente(uc_id=uc.id, tipo_ambiente_id=tid))
        por_matriz.setdefault(mat.id, {})[uc.codigo.lower()] = uc.id
        prereqs += [(mat.id, uc.id, cod) for cod in ext.get("_prereq") or []]
    db.flush()
    from app.models import UCPrerequisito
    for mid, uc_id, cod in prereqs:
        pre = por_matriz.get(mid, {}).get(cod.lower()) or db.scalar(select(UnidadeCurricular.id).where(
            UnidadeCurricular.matriz_id == mid, func.lower(UnidadeCurricular.codigo) == cod.lower()))
        if pre and pre != uc_id and not db.get(UCPrerequisito, (uc_id, pre)):
            db.add(UCPrerequisito(uc_id=uc_id, prerequisito_id=pre))
        elif not pre:
            rel.setdefault("avisos", []).append(f"pré-requisito '{cod}' não encontrado na mesma matriz")
    rel.setdefault("avisos", [])
    rel["avisos"] += [f"linha {it['linha']}: {a}" for it in itens if it["acao"] != "erro" for a in it["avisos"]]
    return rel


def _gravar_levantamento(db, instrutor, ext: dict):
    """Disponibilidade (substitui a anterior), competências e certificações do levantamento."""
    from app.models import HabilitacaoInstrutor
    db.flush()
    from app.services.notificacoes import notificar
    notificar(db, instrutor.id, "cadastro", "Seu cadastro foi atualizado",
              "A coordenação importou o levantamento de instrutores: áreas, disponibilidade, UCs e certificações podem ter mudado.", "dados")
    if ext.get("_disp"):
        db.execute(DisponibilidadeInstrutor.__table__.delete().where(DisponibilidadeInstrutor.instrutor_id == instrutor.id))
        for dia, ini, fim in ext["_disp"]:
            db.add(DisponibilidadeInstrutor(instrutor_id=instrutor.id, dia_semana=dia, hora_inicio=time.fromisoformat(ini), hora_fim=time.fromisoformat(fim)))
    for uc_id, nivel in ext.get("_comp") or []:
        c = db.scalars(select(CompetenciaInstrutor).where(CompetenciaInstrutor.instrutor_id == instrutor.id, CompetenciaInstrutor.uc_id == uc_id)).first()
        if c is None:
            db.add(CompetenciaInstrutor(instrutor_id=instrutor.id, uc_id=uc_id, nivel=nivel, ativo=True))
        else:
            c.nivel, c.ativo = nivel, True
    for nome in ext.get("_cert") or []:
        if not db.scalars(select(HabilitacaoInstrutor).where(HabilitacaoInstrutor.instrutor_id == instrutor.id,
                                                             func.lower(HabilitacaoInstrutor.nome) == nome.lower())).first():
            db.add(HabilitacaoInstrutor(instrutor_id=instrutor.id, nome=nome, ativo=True))


def _confirmar_aulas(db: Session, ctx, imp: Importacao, itens: list[dict]) -> dict:
    from app.services.alocacao import DadosAula, criar_aula, criar_recorrente, obter_cenario_editavel
    if not imp.cenario_id:
        raise EntradaInvalida("Selecione o cenário (em elaboração) que receberá as aulas importadas.")
    obter_cenario_editavel(db, imp.cenario_id)  # cenários publicados nunca são sobrescritos
    criados, sem_alt, bloqueadas, avisos = 0, 0, [], []
    for it in itens:
        d = it["dados"]
        hi, hf = time.fromisoformat(d["hora_inicio"]), time.fromisoformat(d["hora_fim"])
        if d["data"]:
            dt = date.fromisoformat(d["data"])
            existe = db.scalar(select(func.count()).select_from(Alocacao).where(
                Alocacao.cenario_id == imp.cenario_id, Alocacao.turma_uc_id == d["turma_uc_id"], Alocacao.data == dt,
                Alocacao.inicio == datetime.combine(dt, hi), Alocacao.status != "cancelada"))
            if existe:
                sem_alt += 1
                continue
            sp = db.begin_nested()
            try:
                criar_aula(db, ctx, imp.cenario_id, DadosAula(turma_uc_id=d["turma_uc_id"], data=dt, hora_inicio=hi, hora_fim=hf, grupo=d["grupo"],
                                                              instrutor_id=d["instrutor_id"], ambiente_id=d["ambiente_id"], origem="importacao",
                                                              observacao=f"importação #{imp.id} linha {it['linha']}"))
                sp.commit()
                criados += 1
            except ConflitoValidacao as e:
                sp.rollback()
                bloqueadas.append({"linha": it["linha"], "motivo": e.mensagem, "ocorrencias": (e.detalhes or {}).get("ocorrencias", [])[:5]})
        else:
            p = db.scalars(select(PeriodoHorario).where(PeriodoHorario.unidade_id == d["unidade_id"], PeriodoHorario.hora_inicio == hi)).first()
            if not p:
                bloqueadas.append({"linha": it["linha"], "motivo": f"Nenhum período da unidade inicia às {hi:%H:%M}."})
                continue
            seq = db.scalars(select(PeriodoHorario).where(PeriodoHorario.turno_id == p.turno_id, PeriodoHorario.ordem >= p.ordem).order_by(PeriodoHorario.ordem)).all()
            n = next((k + 1 for k, q in enumerate(seq) if q.hora_fim >= hf), None)
            if n is None or seq[n - 1].hora_fim != hf:
                bloqueadas.append({"linha": it["linha"], "motivo": f"O término {hf:%H:%M} não coincide com o fim de um período do turno."})
                continue
            turma = db.get(Turma, d["turma_id"])
            ja = db.scalar(select(func.count()).select_from(Alocacao).where(
                Alocacao.cenario_id == imp.cenario_id, Alocacao.turma_uc_id == d["turma_uc_id"], Alocacao.status != "cancelada",
                func.extract("isodow", Alocacao.data) == d["dia_semana"] + 1, func.to_char(Alocacao.inicio, "HH24:MI") == hi.strftime("%H:%M")))
            if ja:
                sem_alt += 1
                continue
            sp = db.begin_nested()
            try:
                _atv, criadas, _p, _o = criar_recorrente(db, ctx, imp.cenario_id, DadosAula(
                    turma_uc_id=d["turma_uc_id"], data=turma.data_inicio, periodo_inicio_id=p.id, num_periodos=n, grupo=d["grupo"],
                    instrutor_id=d["instrutor_id"], ambiente_id=d["ambiente_id"], origem="importacao"), d["dia_semana"], turma.data_inicio, turma.data_fim,
                    origem="importacao")
                sp.commit()
                criados += len(criadas)
            except ConflitoValidacao as e:
                sp.rollback()
                bloqueadas.append({"linha": it["linha"], "motivo": e.mensagem, "detalhes": str(e.detalhes)[:500] if e.detalhes else None})
    return {"criados": criados, "sem_alteracao": sem_alt, "aulas_bloqueadas": bloqueadas, "avisos": avisos}
